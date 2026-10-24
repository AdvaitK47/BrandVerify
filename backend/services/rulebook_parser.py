from pathlib import Path
import re

def extract_rulebook_text(file_path: str) -> str:
    path = Path(file_path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        import fitz
        with fitz.open(file_path) as doc:
            return "\n".join(page.get_text("text") for page in doc)
    if ext == ".docx":
        from docx import Document
        doc = Document(file_path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    if ext == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError("Supported rulebooks: PDF, DOCX, TXT.")

def draft_rules_from_text(text: str, brand_id: str, filename: str) -> dict:
    colors = sorted(set(re.findall(r"#[0-9A-Fa-f]{6}\b", text)))
    required, prohibited, generic = [], [], []
    for raw in text.splitlines():
        line = re.sub(r"^\s*(?:[-*•]|\d+[.)])\s*", "", raw).strip()
        if not line:
            continue
        low = line.lower()
        if low.startswith(("required phrase:", "required disclaimer:", "must include:")):
            value = line.split(":", 1)[1].strip().strip('"“”')
            if value: required.append(value)
        elif low.startswith(("prohibited phrase:", "banned phrase:", "must not say:")):
            value = line.split(":", 1)[1].strip().strip('"“”')
            if value: prohibited.append(value)
        elif any(x in low for x in ("tone", "logo", "privacy", "claim", "font", "typography", "must not", "do not", "never", "always")):
            generic.append(line)
    return {
        "brand_id": brand_id, "source_filename": filename, "source_text": text,
        "approved_colors": colors, "required_phrases": list(dict.fromkeys(required)),
        "prohibited_phrases": list(dict.fromkeys(prohibited)),
        "generic_rules_for_human_review": list(dict.fromkeys(generic)),
        "approved": False,
        "note": "Draft only. Admin must review and approve before scanning. Only clearly labelled phrases become exact checks."
    }
