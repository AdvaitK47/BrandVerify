"""
Document text extraction (.pdf, .docx).

Turns an uploaded document into plain text so the same claim rules used for
text scans can run on it. Extraction only: no compliance decisions here.
"""

from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional, Union

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


class DocumentError(ValueError):
    """The document is unsupported, unreadable, or cannot be processed."""


@dataclass
class ExtractedDocument:
    # One entry per PDF page. A .docx has no fixed pages, so it is one entry.
    pages: List[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(self.pages)

    @property
    def page_count(self) -> int:
        return len(self.pages)

    @property
    def has_text(self) -> bool:
        return bool(self.text.strip())


def extract_document(
    path: Union[str, Path], filename: Optional[str] = None
) -> ExtractedDocument:
    """Extract text from a .pdf or .docx file on disk."""
    extension = Path(filename or str(path)).suffix.lower()

    if extension == ".pdf":
        return _extract_pdf(path)
    if extension == ".docx":
        return _extract_docx(path)

    raise DocumentError(
        f"Unsupported document type '{extension or 'unknown'}'. "
        f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}."
    )


def _extract_pdf(path: Union[str, Path]) -> ExtractedDocument:
    try:
        import fitz  # PyMuPDF; imported here so a missing install cannot crash the app
    except ImportError as exc:
        raise DocumentError(
            "PDF support needs PyMuPDF. Install it with: pip install PyMuPDF"
        ) from exc

    try:
        with fitz.open(str(path)) as doc:
            if doc.needs_pass:
                raise DocumentError("This PDF is password-protected.")
            pages = [page.get_text() for page in doc]
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError(
            "Could not read this PDF. It may be corrupt."
        ) from exc

    return ExtractedDocument(pages=pages)


def _extract_docx(path: Union[str, Path]) -> ExtractedDocument:
    try:
        from docx import Document  # python-docx
    except ImportError as exc:
        raise DocumentError(
            "Word support needs python-docx. Install it with: pip install python-docx"
        ) from exc

    try:
        document = Document(str(path))
    except Exception as exc:
        raise DocumentError(
            "Could not read this Word file. It may be corrupt."
        ) from exc

    parts: List[str] = [p.text for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    # Legal disclaimers very often live in headers and footers.
    for section in document.sections:
        for block in (section.header, section.footer):
            parts.extend(p.text for p in block.paragraphs if p.text.strip())

    return ExtractedDocument(pages=["\n".join(parts)])
